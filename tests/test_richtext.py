from datetime import date
from io import BytesIO

import pytest
from conftest import pdf_text
from docx import Document
from factories import base_profile, experience
from pypdf import PdfReader

from aptum.common.richtext import (
    Run,
    parse,
    split_sentences,
    to_markdown,
    to_plain,
    to_reportlab,
)
from aptum.modules.cv.ats import rules
from aptum.modules.cv.ats.builder import build_ats_document
from aptum.modules.cv.ats.fidelity import enforce_fidelity, profile_facts
from aptum.modules.cv.ats.pipeline import generate_ats_cv
from aptum.modules.cv.ats.years import years_of_experience
from aptum.modules.cv.document import build_cv_data
from aptum.modules.cv.docx_render import render_cv_docx
from aptum.modules.cv.templates.software_engineer import SoftwareEngineerTemplate
from aptum.modules.profile.schemas import ExperienceCreate, ExperienceUpdate

TODAY = date(2026, 10, 1)
BULLET = "Cut **p95 latency** by *40%* with ***Redis*** caching."
BULLET_PLAIN = "Cut p95 latency by 40% with Redis caching."


# --- Parser ----------------------------------------------------------------------------------


def _styled(text: str) -> list[tuple[str, bool, bool]]:
    """Non-space runs and their style: what a reader sees, ignoring how whitespace is styled."""
    return [(run.text.strip(), run.bold, run.italic) for run in parse(text) if run.text.strip()]


def test_bold_italic_and_both():
    assert parse(BULLET) == [
        Run("Cut "),
        Run("p95 latency", bold=True),
        Run(" by "),
        Run("40%", italic=True),
        Run(" with "),
        Run("Redis", bold=True, italic=True),
        Run(" caching."),
    ]


@pytest.mark.parametrize(
    ("text", "expected"),
    [
        ("**a *b* c**", [("a", True, False), ("b", True, True), ("c", True, False)]),
        ("*a **b** c*", [("a", False, True), ("b", True, True), ("c", False, True)]),
        ("***a** b*", [("a", True, True), ("b", False, True)]),
        ("**a***b*", [("a", True, False), ("b", False, True)]),
    ],
)
def test_nested_and_overlapping_markers(text, expected):
    assert _styled(text) == expected


@pytest.mark.parametrize(
    "text",
    ["**open", "close**", "2 * 3 * 4", "** spaced **", "x ****y**** z", "a*", "[link](https://x.dev) <b>html</b> # h"],
)
def test_unmatched_or_unsupported_syntax_is_literal(text):
    assert parse(text) == [Run(text)]


def test_escapes_are_literal_characters():
    assert parse(r"5 \* 3 \\ done") == [Run(r"5 * 3 \ done")]
    assert parse(r"\*\*not bold\*\*") == [Run("**not bold**")]


@pytest.mark.parametrize(
    "text", [BULLET, "**a *b* c**", "***a** b*", "**open", r"5 \* 3 \\ x", "a **b. c** d.", "plain"]
)
def test_markdown_round_trip_keeps_text_and_style(text):
    again = to_markdown(parse(text))
    assert to_plain(again) == to_plain(text)
    assert _styled(again) == _styled(text)


def test_reportlab_markup_escapes_text_and_keeps_line_breaks():
    assert to_reportlab("a <b> & **c**\n*d*") == "a &lt;b&gt; &amp; <b>c</b><br/><i>d</i>"


def test_split_sentences_closes_formatting_in_each_piece():
    assert split_sentences("Did **x. Then y** done. End") == ["Did **x.**", "**Then y** done.", "End"]


# --- API input ---------------------------------------------------------------------------------


def test_bullets_are_single_line_and_keep_markers():
    data = ExperienceCreate(
        position="Engineer", employer_id=1, start_date="2024-01", functions=["Built **API**\nfor *payments*"]
    )
    assert data.functions == ["Built **API** for *payments*"]
    assert ExperienceUpdate(functions=["a\r\nb"]).functions == ["a b"]


# --- Renderers -------------------------------------------------------------------------------


def _profile():
    profile = base_profile(summary="Engineer with **backend** focus.\nBuilds *LLM* features.")
    profile.experiences[0].description = "Payments **platform**."
    profile.experiences[0].functions[0].description = BULLET
    profile.projects[0].description = "Open **CLI** for *cloud* costs."
    return profile


def _fonts(pdf: bytes) -> set[str]:
    return {
        str(font.get_object()["/BaseFont"])
        for page in PdfReader(BytesIO(pdf)).pages
        for font in page["/Resources"]["/Font"].values()
    }


@pytest.mark.parametrize(
    "render",
    [
        lambda profile: generate_ats_cv(profile, today=TODAY).pdf,
        lambda profile: SoftwareEngineerTemplate().render(build_cv_data(profile)),
    ],
    ids=["basic", "software-engineer"],
)
def test_pdf_prints_formatting_not_markers(render):
    pdf = render(_profile())
    text = " ".join(pdf_text(pdf).split())
    assert "*" not in text
    assert BULLET_PLAIN in text
    assert "Open CLI for cloud costs." in text
    assert {"/Helvetica-Oblique", "/Helvetica-BoldOblique"} <= _fonts(pdf)


def test_docx_bullet_runs_carry_bold_and_italic():
    document = Document(BytesIO(render_cv_docx(build_cv_data(_profile()))))
    bullet = next(p for p in document.paragraphs if p.text == BULLET_PLAIN)
    styled = {(run.text, bool(run.bold), bool(run.italic)) for run in bullet.runs}
    assert {("p95 latency", True, False), ("40%", False, True), ("Redis", True, True)} <= styled
    summary = next(p for p in document.paragraphs if p.text.startswith("Engineer with"))
    assert "*" not in summary.text
    assert summary.text == "Engineer with backend focus.\nBuilds LLM features."


# --- ATS checks see plain text -------------------------------------------------------------------


def test_ats_text_blocks_have_no_markers():
    document = build_ats_document(_profile())
    assert "*" not in document.plain_text()
    assert BULLET_PLAIN in document.plain_text()


def test_duplicate_detection_ignores_markers():
    profile = base_profile()
    profile.experiences = [
        experience("Engineer", "Acme", date(2024, 1, 1), None, ["Built the **billing API** for 5 countries."]),
        experience("Engineer", "Globex", date(2022, 1, 1), date(2023, 12, 1), ["Built the billing API for 5 countries."]),
    ]
    document = build_ats_document(profile)
    warnings = []
    rules.remove_duplicates(document, warnings)
    assert [w.code for w in warnings] == ["duplicate_removed"]
    assert document.experiences[1].bullets == []


def test_fidelity_repair_leaves_no_dangling_markers():
    profile = base_profile()
    years = years_of_experience(profile.experiences, TODAY)
    document = build_ats_document(profile)
    document.summary = "Engineer with **backend focus. Served 987 banks.** Builds tools."
    enforce_fidelity(document, profile, profile_facts(profile, years), [])
    assert document.summary == "Engineer with **backend focus.** Builds tools."
