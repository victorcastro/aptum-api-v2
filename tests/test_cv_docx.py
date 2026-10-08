from io import BytesIO

from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from factories import base_profile

from aptum.modules.cv.document import build_cv_data
from aptum.modules.cv.docx_render import render_cv_docx

CERT_URL = "https://verify.example.dev/abc123"


def _render(profile):
    return Document(BytesIO(render_cv_docx(build_cv_data(profile))))


def _texts(document) -> list[str]:
    return [p.text for p in document.paragraphs]


def _hyperlinks(document) -> list[str]:
    rels = document.part.rels
    return [
        rels[h.get(qn("r:id"))].target_ref
        for p in document.paragraphs
        for h in p._p.findall(qn("w:hyperlink"))
    ]


def test_uses_only_native_structures():
    document = _render(base_profile())
    assert document.tables == []
    assert len(document.inline_shapes) == 0
    assert all(text.strip() for text in _texts(document))


def test_header_and_section_order():
    document = _render(base_profile())
    texts = _texts(document)
    assert texts[:3] == [
        "Alex Rivera",
        "Senior Software Engineer",
        "Toronto, Ontario, Canada | +1 555 0100 | alex.rivera@example.com",
    ]
    headings = [p.text for p in document.paragraphs if p.style.name == "Heading 1"]
    assert headings == [
        "SUMMARY",
        "TECHNICAL SKILLS",
        "EXPERIENCE",
        "PROJECTS",
        "EDUCATION",
        "CERTIFICATIONS",
        "LANGUAGES",
    ]


def test_dates_sit_after_a_tab_with_a_right_tab_stop():
    document = _render(base_profile())
    dated = [p for p in document.paragraphs if p.style.name == "CV Item" and "\t" in p.text]
    assert dated
    for paragraph in dated:
        title, dates = paragraph.text.split("\t")
        assert title and dates
        stops = paragraph.paragraph_format.tab_stops
        assert [s.alignment for s in stops] == [WD_TAB_ALIGNMENT.RIGHT]
        assert abs(stops[0].position - 18 * 360000) < 1000


def test_function_bullets_are_real_list_items():
    document = _render(base_profile())
    bullets = [p for p in document.paragraphs if p.style.name == "List Bullet"]
    assert any("RAG assistant" in p.text for p in bullets)
    assert not any(p.text.startswith("•") for p in document.paragraphs)


def test_institution_and_issuer_are_their_own_paragraph_below_the_title():
    texts = _texts(_render(base_profile()))
    degree = next(i for i, t in enumerate(texts) if t.startswith("BSc, Computer Science"))
    assert "Lakeside University" not in texts[degree]
    assert texts[degree + 1] == "Lakeside University"
    cert = next(i for i, t in enumerate(texts) if t.startswith("Cloud Practitioner"))
    assert texts[cert + 1].startswith("Example Cloud Institute")


def test_header_links_and_certification_link_are_real_hyperlinks():
    profile = base_profile()
    profile.certifications[0].credential_url = CERT_URL
    document = _render(profile)
    links = _hyperlinks(document)
    assert "https://github.com/example-alex" in links
    assert CERT_URL in links
    assert f"Example Cloud Institute | {CERT_URL}" in _texts(document)


def test_certification_link_hidden_when_toggled_off():
    profile = base_profile()
    profile.certifications[0].credential_url = CERT_URL
    profile.certifications[0].show_credential_url = False
    document = _render(profile)
    assert CERT_URL not in _hyperlinks(document)
    assert not any("verify.example.dev" in t for t in _texts(document))


def test_project_link_is_a_hyperlink_and_can_be_hidden():
    profile = base_profile()
    profile.projects[0].url = "https://github.com/example-alex/cli"
    assert "https://github.com/example-alex/cli" in _hyperlinks(_render(profile))
    profile.projects[0].show_url = False
    assert "https://github.com/example-alex/cli" not in _hyperlinks(_render(profile))


def test_inactive_language_is_not_printed():
    profile = base_profile()
    profile.languages[0].is_active = False
    name = profile.languages[0].language.name
    assert not any(name in t for t in _texts(_render(profile)))


def test_ampersand_and_long_titles_render():
    profile = base_profile()
    profile.experiences[0].position = "R&D Engineer & Lead " * 8
    document = _render(profile)
    assert any(t.startswith("R&D Engineer & Lead") for t in _texts(document))


def test_page_is_a4():
    section = _render(base_profile()).sections[0]
    assert round(section.page_width.cm, 1) == 21.0
    assert round(section.page_height.cm, 1) == 29.7
