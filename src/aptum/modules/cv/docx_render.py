from io import BytesIO

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.styles.style import ParagraphStyle
from docx.text.paragraph import Paragraph

from aptum.common.richtext import parse
from aptum.modules.cv.document import CVDocument
from aptum.modules.cv.links import (
    CREDENTIAL_LINK_TEXT,
    format_link_display,
    normalize_link_href,
)

_FONT = "Arial"
_INK = RGBColor(0x1F, 0x29, 0x33)
_ACCENT = RGBColor(0x1F, 0x3A, 0x5F)
_GRAY = RGBColor(0x52, 0x60, 0x6D)
_RULE = "B8C2CC"
_PAGE_WIDTH = Cm(21)
_MARGIN_X = Cm(1.5)
_TEXT_WIDTH = _PAGE_WIDTH - 2 * _MARGIN_X


def render_cv_docx(doc: CVDocument) -> bytes:
    """Editable Word version of the software-engineer layout. Only native features, so Word and
    Google Docs keep it editable: named styles, a right tab stop for dates, real bullet lists and
    hyperlinks. No tables, text boxes, images or empty spacer paragraphs."""
    document = Document()
    _setup_page(document)
    _setup_styles(document)
    document.core_properties.title = f"{doc.full_name} - CV"
    document.core_properties.author = doc.full_name

    document.add_paragraph(doc.full_name, style="Title")
    if doc.headline:
        document.add_paragraph(doc.headline, style="Subtitle")
    if doc.availability_line:
        document.add_paragraph(doc.availability_line, style="CV Contact")
    if doc.phone or doc.email:
        contact = document.add_paragraph(doc.phone or "", style="CV Contact")
        if doc.email:
            if doc.phone:
                contact.add_run(" | ")
            _add_hyperlink(contact, doc.email, f"mailto:{doc.email}", styled=False)
    if doc.links:
        links = document.add_paragraph(style="CV Contact")
        for index, link in enumerate(doc.links):
            if index:
                links.add_run(" | ")
            _add_hyperlink(links, link.label, link.url)

    if doc.summary:
        _heading(document, "Summary")
        _add_rich(document, doc.summary)

    if doc.skill_groups:
        _heading(document, "Technical Skills")
        for group in doc.skill_groups:
            _labeled(document, group.label, ", ".join(group.names))

    if doc.experiences:
        _heading(document, "Experience")
        for exp in doc.experiences:
            _dated(document, exp.title, exp.dates)
            if exp.description:
                _add_rich(document, exp.description)
            for text in exp.bullets:
                _add_rich(document, text, style="List Bullet")
            if exp.skills_line:
                _labeled(document, "Technologies", exp.skills_line, style="CV Muted")

    if doc.projects:
        _heading(document, "Projects")
        for index, project in enumerate(doc.projects):
            name = document.add_paragraph(project.name, style="CV Item")
            if index:
                name.paragraph_format.space_before = Pt(10)
            if project.description:
                _add_rich(document, project.description)
            if project.url:
                _add_url(document.add_paragraph(), project.url)

    if doc.educations:
        _heading(document, "Education")
        for edu in doc.educations:
            _dated(document, edu.title, edu.dates)
            document.add_paragraph(edu.institution)

    if doc.certifications:
        _heading(document, "Certifications")
        for index, cert in enumerate(doc.certifications):
            title = _dated(document, cert.title, cert.dates)
            if index:
                title.paragraph_format.space_before = Pt(10)
            line = document.add_paragraph(cert.issuer)
            if cert.url:
                line.add_run(" | ")
                _add_hyperlink(line, CREDENTIAL_LINK_TEXT, normalize_link_href(cert.url))

    if doc.languages:
        _heading(document, "Languages")
        for language in doc.languages:
            document.add_paragraph(language)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _setup_page(document) -> None:
    section = document.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = _PAGE_WIDTH
    section.page_height = Cm(29.7)
    section.left_margin = section.right_margin = _MARGIN_X
    section.top_margin = section.bottom_margin = Cm(1)


def _set_font(style: ParagraphStyle, size: float, bold: bool = False, color: RGBColor = _INK) -> None:
    font = style.font
    font.name = _FONT
    font.size = Pt(size)
    font.bold = bold
    font.italic = False
    font.color.rgb = color
    rfonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        rfonts.attrib.pop(qn(f"w:{attr}"), None)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), _FONT)


def _clear_border(style: ParagraphStyle) -> None:
    ppr = style.element.pPr
    if ppr is not None:
        for border in ppr.findall(qn("w:pBdr")):
            ppr.remove(border)


def _set_bottom_border(style: ParagraphStyle) -> None:
    _clear_border(style)
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, value in (("val", "single"), ("sz", "4"), ("space", "1"), ("color", _RULE)):
        bottom.set(qn(f"w:{key}"), value)
    border.append(bottom)
    style.element.get_or_add_pPr().append(border)


def _custom_style(styles, name: str, base: str = "Normal") -> ParagraphStyle:
    style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = styles[base]
    style.quick_style = True
    return style


def _setup_styles(document) -> None:
    styles = document.styles

    normal = styles["Normal"]
    _set_font(normal, 10)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(2)
    normal.paragraph_format.line_spacing = 1.15

    title = styles["Title"]
    _set_font(title, 21, bold=True, color=_ACCENT)
    _clear_border(title)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(2)

    subtitle = styles["Subtitle"]
    _set_font(subtitle, 11.5, color=_GRAY)
    subtitle.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(4)

    contact = _custom_style(styles, "CV Contact")
    _set_font(contact, 9.5, color=_GRAY)
    contact.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    muted = _custom_style(styles, "CV Muted")
    _set_font(muted, 9, color=_GRAY)

    heading = styles["Heading 1"]
    _set_font(heading, 10.5, bold=True, color=_ACCENT)
    heading.paragraph_format.space_before = Pt(14)
    heading.paragraph_format.space_after = Pt(4)
    heading.paragraph_format.keep_with_next = True
    _set_bottom_border(heading)

    item = _custom_style(styles, "CV Item")
    _set_font(item, 10, bold=True)
    item.paragraph_format.space_before = Pt(6)
    item.paragraph_format.space_after = Pt(0)
    item.paragraph_format.keep_with_next = True
    item.paragraph_format.tab_stops.add_tab_stop(_TEXT_WIDTH, WD_TAB_ALIGNMENT.RIGHT)

    bullet = styles["List Bullet"]
    _set_font(bullet, 10)
    bullet.paragraph_format.left_indent = Cm(0.6)
    bullet.paragraph_format.first_line_indent = Cm(-0.4)
    bullet.paragraph_format.space_after = Pt(1)


def _heading(document, text: str) -> None:
    document.add_paragraph(text.upper(), style="Heading 1")


def _add_rich(document, text: str, style: str | None = None) -> Paragraph:
    """User free text with its bold and italic (see `common/richtext.py`); line breaks stay in the paragraph."""
    paragraph = document.add_paragraph(style=style)
    for run in parse(text):
        for index, line in enumerate(run.text.split("\n")):
            if index:
                paragraph.add_run().add_break()
            if line:
                added = paragraph.add_run(line)
                added.bold = run.bold or None
                added.italic = run.italic or None
    return paragraph


def _labeled(document, label: str, value: str, style: str | None = None) -> None:
    paragraph = document.add_paragraph(style=style)
    paragraph.add_run(f"{label}: ").bold = True
    paragraph.add_run(value)


def _dated(document, title: str, dates: str) -> Paragraph:
    """Title, then a tab to the right margin and the date, all in one paragraph."""
    paragraph = document.add_paragraph(style="CV Item")
    paragraph.paragraph_format.tab_stops.add_tab_stop(_TEXT_WIDTH, WD_TAB_ALIGNMENT.RIGHT)
    paragraph.add_run(title)
    if dates:
        run = paragraph.add_run(f"\t{dates}")
        run.bold = False
        run.font.size = Pt(9)
        run.font.color.rgb = _GRAY
    return paragraph


def _add_url(paragraph: Paragraph, url: str) -> None:
    """The clean URL as text (no scheme, no `www.`), linked to the full `https://` URL (project links)."""
    _add_hyperlink(paragraph, format_link_display(url), normalize_link_href(url))


def _add_hyperlink(paragraph: Paragraph, text: str, url: str, *, styled: bool = True) -> None:
    """`styled=False` keeps the paragraph's own look (the email on the contact line)."""
    relation_id = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation_id)
    run = OxmlElement("w:r")
    if styled:
        props = OxmlElement("w:rPr")
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "1F3A5F")
        underline = OxmlElement("w:u")
        underline.set(qn("w:val"), "single")
        props.append(color)
        props.append(underline)
        run.append(props)
    content = OxmlElement("w:t")
    content.text = text
    content.set(qn("xml:space"), "preserve")
    run.append(content)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
